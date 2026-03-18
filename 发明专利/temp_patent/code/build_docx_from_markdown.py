# -*- coding: utf-8 -*-
import os
import sys

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = document.styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)


def add_paragraph(document, text):
    para = document.add_paragraph()
    para.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    para.paragraph_format.first_line_indent = Pt(24)
    para.add_run(text)


def build(markdown_path, docx_path):
    document = Document()
    configure_styles(document)

    with open(markdown_path, "r", encoding="utf-8") as fp:
        lines = [line.rstrip() for line in fp]

    for line in lines:
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith("1. ") or line.startswith("2. ") or line.startswith("3. ") or line.startswith("4. ") or line.startswith("5. "):
            document.add_paragraph(line.strip(), style="List Number")
        else:
            add_paragraph(document, line)

    document.save(docx_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python build_docx_from_markdown.py <input.md> <output.docx>")
    md_path = os.path.abspath(sys.argv[1])
    docx_path = os.path.abspath(sys.argv[2])
    build(md_path, docx_path)
