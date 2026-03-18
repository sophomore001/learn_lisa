# -*- coding: utf-8 -*-
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


def delete_paragraph(paragraph):
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def insert_paragraph_after(paragraph, text):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if paragraph.style is not None:
        try:
            new_para.style = paragraph.style
        except Exception:
            pass
    run = new_para.add_run(text)
    font = run.font
    font.name = "宋体"
    return new_para


def replace_paragraph_text(paragraph, text):
    for run in paragraph.runs[::-1]:
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)


def find_paragraph_by_prefix(cell, prefix):
    for paragraph in cell.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    return None


def find_main_docx(base_dir):
    candidates = [
        n
        for n in os.listdir(base_dir)
        if n.startswith("202602191048")
        and n.endswith(".docx")
        and "公式" not in n
        and "修订" not in n
    ]
    if not candidates:
        raise FileNotFoundError("未找到原始交底书 docx")
    return os.path.join(base_dir, candidates[0])


def revise_document(input_path, output_path):
    doc = Document(input_path)

    # 删除首页模板提示段落
    for idx in [6, 5, 4, 3, 2, 1, 0]:
        delete_paragraph(doc.paragraphs[idx])

    table = doc.tables[1]

    # 二、背景技术
    bg_para = table.rows[3].cells[0].paragraphs[0]
    replace_paragraph_text(
        bg_para,
        "星载超稳晶振（Ultra-Stable Oscillator，USO）作为航天器时间频率基准的重要组成部分，其频率稳定性直接影响深空探测、星间链路同步和空间引力波探测等任务的测量精度。公开文献表明，石英晶体振荡器的静态频率—温度特性通常可用多项式模型描述，其中三次模型较为常见；但在温度循环或快变温条件下，升温与降温路径会出现不重合现象，即热滞后现象（Xiaogang Deng 等，IEEE TUFFC，2021）。Wang 和 Wu（Sensors，2020）进一步指出，传统 TCXO 往往只依据当前温度进行补偿，而热滞后补偿需要同时考虑当前温度信息和温度变化历史信息，并且温度传感器与晶体之间存在 thermal lag。IEEE Std 1193-2022 也指出，频率标准的环境敏感度测量会受到非线性、不同时间常数、瞬态效应和滞后的影响，因此在动态变温环境下仅依赖静态频温曲线难以充分描述器件实际响应。基于此，现有静态温度补偿方案在复杂非稳态热环境下仍存在参数辨识不充分、动态热滞后刻画不足以及中长时间稳定度改善有限等问题，需要引入兼顾静态频温骨架和动态热惯性的补偿模型。"
    )
    replace_paragraph_text(
        table.rows[3].cells[0].paragraphs[1],
        "现有温度补偿方法主要包括恒温控制和静态查表/多项式补偿两类。前者受体积、重量和功耗约束，后者通常假设外部测温点与内部谐振器实时处于热平衡状态。当器件存在隔热封装、热阻和热容耦合时，上述假设在快变温或热循环条件下难以严格成立，因此仍有必要建立能够反映热惯性与热滞后的动态补偿模型。"
    )

    # 三、技术问题，修正过强表述
    tech_p1 = table.rows[5].cells[0].paragraphs[1]
    replace_paragraph_text(
        tech_p1,
        "此外，USO 输出频率中通常叠加白频率噪声、闪烁频率噪声和随机游走频率噪声等随机项，使得温度引起的确定性频率漂移在低频段容易与噪声项相互混杂，从而影响模型参数的稳定估计。对于本发明而言，需要进一步解决的问题是：如何利用历史温度数据及其同步频率数据，在含噪条件下较稳健地估计静态温度系数和热时间常数，并在在线阶段利用实时温度数据和上一时刻模型状态量实现动态温度补偿，从而改善中长积分时间上的频率稳定度。"
    )

    # 四、技术方案
    row7 = table.rows[7].cells[0]
    p1 = row7.paragraphs[1]
    replace_paragraph_text(
        p1,
        "本发明为解决上述问题，提出一种基于热惯性模型的超稳晶振（Ultra-Stable Oscillator，USO）动态温度补偿方法。该方法以静态多项式频温模型描述基础温漂特性，以一阶热惯性递推模型描述内部谐振器相对于外部测温点的动态滞后响应，再利用历史温度与同步频率数据对模型参数进行辨识，最后在在线阶段根据实时温度和上一时刻状态量递推计算动态温度频偏，并对实时输出频率进行修正。其中，静态频温多项式可参考 Deng 等（IEEE TUFFC，2021）和 Haapala 等（IEEE TCAS-I，2020）关于多项式频温建模的研究；热滞后与温度历史信息的必要性可参考 Wang 和 Wu（Sensors，2020）；温度敏感度、热时间常数以及相关分析方法可参考 IEEE Std 1193-2022。"
    )

    p10 = row7.paragraphs[10]
    replace_paragraph_text(
        p10,
        "S4、选取一段历史温度数据及其同步采集的原始频率数据作为训练集，采用非线性最小二乘法进行参数辨识，以模型输出频率与实测频率之间的残差平方和最小为目标，联合估计步骤 S2 与步骤 S3 中涉及的参数；"
    )

    s4a = insert_paragraph_after(
        p10,
        "优选地，步骤 S4 包括：S41、从训练集中筛选温度变化速率较小的准稳态样本段，用于获得静态温度系数初值；S42、利用升温/降温斜坡段、温度阶跃段或热循环段获得热时间常数初值；S43、以前述初值为起点，对全部训练样本执行联合非线性最小二乘精修。该做法有利于降低静态温度系数与热时间常数之间的参数代偿。"
    )
    s4b = insert_paragraph_after(
        s4a,
        "令待估参数向量为 θ = [α, β, γ, τ_th]^T，依据公式（1）至公式（3）构成前向模型，递推得到模型估计频率偏差 ŷ(k; θ)，则可构造目标函数 J(θ) = (1/N)Σ[k=1..N](f_m(k) - ŷ(k; θ))^2，并通过非线性最小二乘法求得使 J(θ) 最小的参数估计结果。对于静态多项式部分，其初值可通过最小二乘拟合获得；Haapala 等（2020）亦指出，多项式温补模型系数可通过 ordinary least squares estimation 重新计算。"
    )
    s4c = insert_paragraph_after(
        s4b,
        "依据 IEEE Std 1193-2022 Annex A，在进行静态温度敏感度估计时，优选先识别器件热时间常数并剔除若干个 1/e 时间常数内的非平衡数据；在低信噪比条件下，还可先进行频率与温度的相关分析，以判断温度敏感性是否具有统计显著性，再实施参数联合拟合。"
    )

    # 五、有益效果与模型评估
    row9 = table.rows[9].cells[0]
    replace_paragraph_text(
        row9.paragraphs[1],
        "针对航天器运行过程中存在的周期性变温、载荷热扰动和非稳态热传导问题，本发明通过将静态频温关系与动态热惯性机制结合，可在升温、降温及快变温场景下更合理地刻画晶体谐振器的温度响应。与仅依据当前温度进行修正的静态补偿方案相比，本发明不仅考虑了温度的瞬时取值，还考虑了温度变化历史通过状态递推对当前频率偏差的影响，因此更适合用于存在 thermal lag 的应用场景。"
    )
    replace_paragraph_text(
        row9.paragraphs[2],
        "为了验证本方法的有效性，可采用训练集与验证集分离的方式对模型进行评估：训练集用于估计参数，验证集用于独立评价模型泛化能力和补偿效果。仿真或试验工况可覆盖典型在轨设备的变温范围与热循环条件，以考察所估参数在非稳态热环境下的适用性。"
    )
    replace_paragraph_text(
        row9.paragraphs[3],
        "图2给出了补偿前后的频率偏差时域对比结果。未补偿情况下，频率偏差随温度变化表现出明显滞后响应；采用本发明方法后，时域残差明显减小，说明由动态热响应引起的确定性频偏得到了有效抑制。"
    )
    replace_paragraph_text(
        row9.paragraphs[6],
        "图3给出了补偿前后的频率—温度回滞环对比结果。未补偿数据在升温与降温过程中形成明显回滞环；经本发明方法补偿后，回滞环面积显著缩小，说明所建立的热惯性模型能够较好反映外部测温点与内部谐振器热状态之间的动态映射关系。"
    )
    replace_paragraph_text(
        row9.paragraphs[9],
        "图4给出了补偿前后的 Allan 偏差对比结果。未补偿数据在中长积分时间段受温度漂移调制而出现稳定度恶化；经补偿后，千秒量级附近的 Allan 偏差明显降低，表明本发明方法对中长期温度相关漂移具有较好的抑制作用。"
    )
    replace_paragraph_text(
        row9.paragraphs[11],
        "图4补偿前后的阿伦偏差对比图"
    )
    replace_paragraph_text(
        row9.paragraphs[12],
        "综上，通过对训练集进行参数辨识并在验证数据上评价 RMSE、温频相关性、回滞环面积及 Allan 偏差，可表明本发明方法在动态变温场景下能够更准确地描述并补偿由热惯性引起的频率漂移，从而提升超稳晶振输出频率的稳定性。"
    )

    eval_p1 = insert_paragraph_after(
        row9.paragraphs[2],
        "优选地，模型准确性至少可由以下指标进行评价：其一，频率残差均方根误差 RMSE = sqrt[(1/N)Σ(e(k)^2)]，其中 e(k) = f_m(k) - ŷ(k)；其二，温度与频率残差之间的相关系数 r_Te，用于衡量补偿后温度相关性是否下降；其三，频率—温度回滞环面积 A_h = ∮ f dT，用于衡量升温与降温路径差异是否收缩；其四，补偿前后 Allan 偏差 σ_y(τ) 在百秒至万秒积分时间范围内的变化，用于评价中长期稳定度是否改善。"
    )
    insert_paragraph_after(
        eval_p1,
        "其中，IEEE Std 1193-2022 Annex A 指出，在低信噪比条件下可通过频率—温度相关分析判断温度敏感性是否显著；同一标准还强调应区分确定性环境漂移与随机频率波动，避免将低频环境效应误判为随机不稳定度。因而，本发明优选同时结合残差误差、相关性、回滞环和 Allan 偏差等多种指标评价模型准确性。"
    )

    # 六、摘要修订
    row11 = table.rows[11].cells[0]
    replace_paragraph_text(
        row11.paragraphs[1],
        "本发明公开了一种基于热惯性模型的超稳晶振（Ultra-Stable Oscillator，USO）动态温度补偿方法，旨在解决非稳态热环境下由于热传导延迟导致传统静态补偿精度不足的问题。该方法首先建立静态多项式频温模型和一阶热惯性递推模型，以描述温度变化对频率偏差的静态和动态影响；然后利用历史温度数据及其同步频率数据，根据公式（1）至公式（3）构造前向模型，并采用非线性最小二乘法联合估计静态温度系数和热时间常数；最后利用实时温度数据和上一时刻模型状态递推动态温度频偏，并按照公式（4）对输出频率进行在线修正。通过对残差误差、温频相关性、回滞环面积及 Allan 偏差进行综合评价，本发明能够有效减小动态变温环境下的温度相关频率漂移。"
    )

    # 删除摘要后模板说明
    for idx in range(len(row11.paragraphs) - 1, 3, -1):
        delete_paragraph(row11.paragraphs[idx])

    # 清理并补齐行文顺序
    row7 = table.rows[7].cells[0]
    for prefix in [
        "上述步骤主要是通过采集环境温度与频率信号",
        "上述步骤的核心在于：",
    ]:
        para = find_paragraph_by_prefix(row7, prefix)
        if para is not None:
            delete_paragraph(para)

    p_formula4 = find_paragraph_by_prefix(row7, "式中，为经本发明方法补偿后的高精度频率输出")
    if p_formula4 is not None:
        online_p = insert_paragraph_after(
            p_formula4,
            "S5 所对应的在线阶段不再重新估计参数，而是将步骤 S4 得到的参数固定，并根据实时温度数据和上一时刻动态状态量递推当前动态温度频偏。其中，历史温度数据主要用于离线辨识，实时温度数据主要用于在线补偿，而上一时刻动态状态量用于保持热惯性模型的短时记忆。"
        )
        insert_paragraph_after(
            online_p,
            "上述步骤的核心在于：利用公式（1）描述静态频温骨架，利用公式（2）和公式（3）描述热惯性导致的动态迟滞，并利用公式（4）完成实时频率修正。由此，本发明将“历史温度 + 同步频率用于参数辨识”与“实时温度 + 上一时刻状态用于在线补偿”统一于同一技术框架中。整个方案计算流程简图如图1所示。"
        )

    row9 = table.rows[9].cells[0]
    for prefix in [
        "图2绘制了补偿前后的频率偏差时域对比图",
        "图3绘制了补偿前后的热滞后回环修正效果",
        "图4绘制了补偿前后的阿伦偏差",
        "综上所述，由仿真结果可知",
    ]:
        para = find_paragraph_by_prefix(row9, prefix)
        if para is not None:
            delete_paragraph(para)

    doc.save(output_path)


if __name__ == "__main__":
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    input_doc = find_main_docx(base_dir)
    output_doc = os.path.join(
        base_dir,
        "202602191048一种基于热惯性模型的超稳晶振动态温度补偿方法_参数估计修订版.docx",
    )
    revise_document(input_doc, output_doc)
    print(output_doc)
